import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_deck():
    template_path = r"C:\Users\Murat\Downloads\Girisim_Sunumu_Sablonu.pptx"
    output_path_1 = r"C:\Users\Murat\Desktop\Trustia\Kurumsal\Sunumlar\Trustia_ASELSAN_Girisim_Sunumu.pptx"
    output_path_2 = r"C:\Users\Murat\Downloads\Trustia_ASELSAN_Girisim_Sunumu.pptx"

    prs = Presentation(template_path)

    # Color Palette
    C_NAVY = RGBColor(11, 37, 69)
    C_BLUE = RGBColor(0, 86, 179)
    C_DARK = RGBColor(33, 37, 41)
    C_GRAY = RGBColor(108, 117, 125)
    C_LIGHT = RGBColor(245, 247, 250)
    C_WHITE = RGBColor(255, 255, 255)
    C_GREEN = RGBColor(40, 167, 69)
    C_BORDER = RGBColor(218, 224, 233)

    logo_path = r"c:\Users\Murat\Desktop\Trustia\Kurumsal\Medya\Logo.png"
    car_path = r"c:\Users\Murat\Desktop\Trustia\Kurumsal\Medya\Arac_Lidar.png"

    def clear_slide(slide):
        for shape in list(slide.shapes):
            sp = shape._element
            sp.getparent().remove(sp)

    def add_header(slide, title, category="ASELSAN GİRİŞİMCİLİK MERKEZİ (AXCELERATE) BAŞVURU SUNUMU"):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_BLUE

        tb2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.65))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
        p_title = tf2.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = C_NAVY

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.7), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = C_BORDER
        line.line.color.rgb = C_BORDER

    # SLIDE 1: KAPAK
    s1 = prs.slides[0]
    clear_slide(s1)
    bg_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.7), Inches(5.9))
    bg_card.fill.solid()
    bg_card.fill.fore_color.rgb = C_LIGHT
    bg_card.line.color.rgb = C_BORDER
    bg_card.line.width = Pt(1.5)

    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(1.3), Inches(1.3), width=Inches(1.8))

    tb = s1.shapes.add_textbox(Inches(1.3), Inches(2.2), Inches(10.5), Inches(2.6))
    tf = tb.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "TRUSTIA TEKNOLOJİ"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = C_BLUE

    p1 = tf.add_paragraph()
    p1.text = "Milli Otonom Sürüş Yazılımı & İnsansız Araç Beyni"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = C_NAVY

    p2 = tf.add_paragraph()
    p2.text = "Uydudan Bağımsız Lazerle Harita Çıkaran ve Sinyal Kesilmesinden Etkilenmeyen Yerli Sistem"
    p2.font.size = Pt(15)
    p2.font.color.rgb = C_GRAY

    meta_box = s1.shapes.add_textbox(Inches(1.3), Inches(4.9), Inches(10.5), Inches(1.4))
    mtf = meta_box.text_frame
    mtf.word_wrap = True
    mp1 = mtf.paragraphs[0]
    mp1.text = "Başvuru: ASELSAN Girişimcilik Merkezi (Axcelerate) & Millileştirme Değerlendirmesi"
    mp1.font.size = Pt(12)
    mp1.font.bold = True
    mp1.font.color.rgb = C_NAVY

    mp2 = mtf.add_paragraph()
    mp2.text = "Sunucu: Murat Furkan Bayram (Kurucu & Sistem Mimarı) | Donanım: Denizcan Özcan (ASELSAN Aday Havuzu)"
    mp2.font.size = Pt(11)
    mp2.font.color.rgb = C_DARK

    mp3 = mtf.add_paragraph()
    mp3.text = "Tarih: 23 Eylül 2026 | Yerleşke: İTO BTM Fulya Kampüsü | Tedarikçi Başvuru No: 0050569"
    mp3.font.size = Pt(11)
    mp3.font.color.rgb = C_GRAY

    # SLIDE 2: PROBLEM
    s2 = prs.slides[1]
    clear_slide(s2)
    add_header(s2, "PROBLEM: Askeri Sahada Yaşanan 3 Büyük Zorluk")
    problems = [
        ("1. Uydu (GPS) Sinyalinin Kesilmesi",
         "Savaş ve operasyon alanlarında sinyal bozucular uydu bağlantısını tamamen keser. Pazardaki standart insansız araçlar uydu bağlantısı koptuğunda yönünü kaybedip durur ve açık hedef haline gelir."),
        ("2. Uzaktan Kumanda (Telsiz) Bağlantısının Kopması",
         "Uzaktan kumandayla yönlendirilen araçlarda telsiz bağı koptuğunda araç olduğu yerde donup kalır. Ayrıca sinyalin gecikmeli gitmesi ani engellerde kazalara yol açar."),
        ("3. Yabancı Otonom Yazılımlara Bağımlılık",
         "Piyasadaki hazır otonom sistemler yabancı kaynaklıdır. Güvenlik açıkları taşır ve Türk savunma sanayii sistemleriyle doğrudan haberleşemez.")
    ]
    for i, (p_title, p_desc) in enumerate(problems):
        top_pos = Inches(1.7 + i * 1.6)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.7), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = C_LIGHT
        card.line.color.rgb = C_BORDER
        tb = s2.shapes.add_textbox(Inches(1.1), top_pos + Inches(0.15), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = p_title
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = RGBColor(180, 40, 40)
        pd = tf.add_paragraph()
        pd.text = p_desc
        pd.font.size = Pt(12)
        pd.font.color.rgb = C_DARK

    # SLIDE 3: ÇÖZÜM
    s3 = prs.slides[2]
    clear_slide(s3)
    add_header(s3, "ÇÖZÜM: Trustia Yerli Otonom Sürüş Çekirdeği")
    solutions = [
        ("1. Uydu Olmadan Lazer Gözlerle (LiDAR) Haritalama",
         "Uydu sinyali tamamen kesilse bile; araç üzerindeki 3 boyutlu lazer tarayıcılar sayesinde araç kendi haritasını arazide kendisi çıkarır ve yolundan şaşmadan ilerler."),
        ("2. Bağlantı Koptuğunda Kendi Başına Üsse Dönüş",
         "Kumanda veya telsiz bağlantısı koptuğu anda araç panik yapmaz. Hafızasındaki lazer haritayı takip ederek güvenli başlangıç noktasına kendi kendine geri döner."),
        ("3. %100 Yerli Kod ve ASELSAN Sistemlerine Tam Uyum",
         "Yazılımın tüm satırları sıfırdan yerli olarak yazılmıştır. Askeri haberleşme kurallarına ve ASELSAN kara sistemlerine doğrudan bağlanabilir.")
    ]
    for i, (s_title, s_desc) in enumerate(solutions):
        top_pos = Inches(1.7 + i * 1.6)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.7), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = C_LIGHT
        card.line.color.rgb = C_BLUE
        card.line.width = Pt(1.5)
        tb = s3.shapes.add_textbox(Inches(1.1), top_pos + Inches(0.15), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = s_title
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = C_NAVY
        pd = tf.add_paragraph()
        pd.text = s_desc
        pd.font.size = Pt(12)
        pd.font.color.rgb = C_DARK

    # SLIDE 4: ÜRÜN
    s4 = prs.slides[3]
    clear_slide(s4)
    add_header(s4, "ÜRÜN: Trustia Otonomi Beyni ve Kontrol Ekranı")
    tb_left = s4.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(6.8), Inches(4.9))
    tfl = tb_left.text_frame
    tfl.word_wrap = True
    features = [
        ("Akıllı Sürüş ve Rota Çizici:", "Lazer sensörlerle çevreyi tarar; yoldaki çukurları, kayaları ve hareketli engelleri anında görüp etrafından dolaşır."),
        ("Askeri Tehlike Algılama Modülü:", "Yoldaki mayın ve patlayıcı şüpheli cisimleri tespit eder, aracı güvenli mesafede durdurur."),
        ("Taktik Komuta Ekranı:", "Askeri harita destekli kontrol ekranı. Operatör ekran üzerinden tek dokunuşla araca hedef verir veya anında durdurur."),
        ("1.301 Test ile Kanıtlanmış Güvenlik:", "16 bin satırlık yazılım çekirdeği 1.301 farklı zorlu arazi senaryosunda test edilmiş ve tamamını hatasız geçmiştir.")
    ]
    for idx, (head, body) in enumerate(features):
        p = tfl.paragraphs[0] if idx == 0 else tfl.add_paragraph()
        p.text = "• " + head + " "
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = C_NAVY
        run = p.add_run()
        run.text = body
        run.font.size = Pt(12)
        run.font.color.rgb = C_DARK

    spec_card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.9), Inches(1.7), Inches(4.6), Inches(4.9))
    spec_card.fill.solid()
    spec_card.fill.fore_color.rgb = C_LIGHT
    spec_card.line.color.rgb = C_BORDER
    if os.path.exists(car_path):
        s4.shapes.add_picture(car_path, Inches(8.1), Inches(1.9), width=Inches(4.2))
    tb_spec = s4.shapes.add_textbox(Inches(8.1), Inches(4.4), Inches(4.2), Inches(2.0))
    tfs = tb_spec.text_frame
    tfs.word_wrap = True
    sp1 = tfs.paragraphs[0]
    sp1.text = "TEKNİK GÜCÜMÜZ"
    sp1.font.size = Pt(12)
    sp1.font.bold = True
    sp1.font.color.rgb = C_BLUE
    specs = [
        "Haberleşme: Askeri ve Robotik Standartlara Tam Uyumlu",
        "Bilgisayar: Araç İçi Güçlü İşlemciler (Nvidia Jetson vb.)",
        "Sensörler: 3 Boyutlu Lazer (LiDAR), Kamera, Açı Ölçerler",
        "Tepki Hızı: 5 Milisaniyede (Göz Kırpmasından Hızlı) Müdahale",
        "Durum: Çalışır Durumda, Göreve Hazır"
    ]
    for s in specs:
        p = tfs.add_paragraph()
        p.text = "✓ " + s
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_DARK

    # SLIDE 5: ALTYAPILAR
    s5 = prs.slides[4]
    clear_slide(s5)
    add_header(s5, "ALTYAPILAR: Çalışma Merkezimiz ve Belgelerimiz")
    infra_boxes = [
        ("İTO BTM Fulya Kampüsü", "İstanbul Ticaret Odası bünyesindeki Fulya Derin Teknoloji Kampüsü'nde yerleşik merkezimiz. Ar-Ge, yazılım ve çalışma alanımız buradadır."),
        ("Yazılım ve Donanım Test Cihazları", "Araç bilgisayarları, 3 boyutlu lazer tarayıcılar (LiDAR), yön algılayıcı sensörler ve araç içi bağlantı ekipmanlarımız mevcuttur."),
        ("1.301 Testli Sanal Test Odası", "Gerçek araziye çıkmadan önce yazılımımızı 1.301 farklı sanal senaryoda (sis, çamur, gece, karıştırma) 43 saniyede otomatik test eden sistemimiz aktiftir."),
        ("Resmi Belgelerimiz ve Onaylarımız", "ASELSAN Tedarikçi Portalı Yazılım Ön Onayı (SAP: FZQHEXGFMTJU), NATO Tedarikçi Kodu Başvurusu, TÜBİTAK ARBİS Kaydı ve Savunma Sanayii Sertifikalarımız.")
    ]
    for idx, (ititle, idesc) in enumerate(infra_boxes):
        row = idx // 2
        col = idx % 2
        l = Inches(0.8 + col * 6.0)
        t = Inches(1.7 + row * 2.5)
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(5.7), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = C_LIGHT
        card.line.color.rgb = C_BORDER
        tb = s5.shapes.add_textbox(l + Inches(0.25), t + Inches(0.2), Inches(5.2), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = ititle
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY
        p2 = tf.add_paragraph()
        p2.text = idesc
        p2.font.size = Pt(12)
        p2.font.color.rgb = C_DARK

    # SLIDE 6: PAZAR
    s6 = prs.slides[5]
    clear_slide(s6)
    add_header(s6, "PAZAR BÜYÜKLÜĞÜ: Askeri İnsansız Araçlar Pazarı")
    market_cards = [
        ("DÜNYA PAZARI", "15.8 Milyar $", "Dünyadaki askeri insansız kara araçları ve taktik otonomi sistemlerinin toplam yıllık pazar büyüklüğü."),
        ("BÖLGEMİZDEKİ PAZAR", "1.2 Milyar $", "Türkiye, NATO ve bölge ülkelerindeki sınır güvenliği, otonom konvoy ve taktik araç yazılımları pazarı."),
        ("BİZİM HEDEFİMİZ", "45 Milyon $", "ASELSAN ve yerli kara araçlarının otonomlaştırılması projelerinde ilk 3 yılda hedeflediğimiz yerli yazılım payı.")
    ]
    for idx, (m_type, m_val, m_desc) in enumerate(market_cards):
        l = Inches(0.8 + idx * 4.0)
        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, Inches(1.8), Inches(3.7), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = C_LIGHT
        card.line.color.rgb = C_BLUE
        card.line.width = Pt(1.5)
        tb = s6.shapes.add_textbox(l + Inches(0.25), Inches(2.1), Inches(3.2), Inches(4.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = m_type
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY
        p2 = tf.add_paragraph()
        p2.text = m_val
        p2.font.size = Pt(26)
        p2.font.bold = True
        p2.font.color.rgb = C_BLUE
        p3 = tf.add_paragraph()
        p3.text = m_desc
        p3.font.size = Pt(12)
        p3.font.color.rgb = C_DARK

    # SLIDE 7: RAKİP ANALİZİ
    s7 = prs.slides[6]
    clear_slide(s7)
    add_header(s7, "FARKIMIZ: Sahadaki Çözümlerle Karşılaştırma")
    rows, cols = 6, 4
    table_shape = s7.shapes.add_table(rows, cols, Inches(0.8), Inches(1.7), Inches(11.7), Inches(4.8))
    table = table_shape.table
    table.columns[0].width = Inches(3.6)
    table.columns[1].width = Inches(2.7)
    table.columns[2].width = Inches(2.7)
    table.columns[3].width = Inches(2.7)
    headers = ["Özellik ve Yetenek", "TRUSTIA (Yerli Yazılım)", "Klasik Kumandalı Araçlar", "Yabancı Açık Kaynaklar"]
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_NAVY
        p = cell.text_frame.paragraphs[0]
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = C_WHITE
        p.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT

    matrix_data = [
        ("Uydusuz Lazerle Harita Çıkarma", "✓ Tam Uyumlu (Lazerle Gider)", "✗ Yok (Uydu Şart)", "△ Şehir Haritası İster"),
        ("Askeri Standartlara Uyum", "✓ Tam Uyumlu", "△ Sadece Telsiz", "✗ Askeri Standart Yok"),
        ("%100 Yerli Kod Güvencesi", "✓ Tamamen Yerli", "△ Parçalara Bağımlı", "✗ Yabancıya Bağımlı"),
        ("Bağlantı Kopunca Eve Dönüş", "✓ Kendi Kendine Döner", "✗ Araç Olduğu Yerde Kalır", "△ Belirsiz"),
        ("1.301 Testten Geçmiş Sistem", "✓ Test Edilmiş & Hazır", "△ Deneme Aşamasında", "△ Sivil Testler")
    ]
    for row_idx, row_data in enumerate(matrix_data, 1):
        for col_idx, text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = text
            cell.fill.solid()
            cell.fill.fore_color.rgb = C_LIGHT if row_idx % 2 == 0 else C_WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            p.font.bold = (col_idx == 1)
            if col_idx == 1:
                p.font.color.rgb = C_GREEN
            elif "✗" in text:
                p.font.color.rgb = RGBColor(200, 30, 30)
            else:
                p.font.color.rgb = C_DARK
            p.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT

    # SLIDE 8: İŞ MODELİ
    s8 = prs.slides[7]
    clear_slide(s8)
    add_header(s8, "İŞ MODELİ: Savunma Odaklı Yazılım Lisanslama")
    models = [
        ("1. Araç Başına Yazılım Lisansı Satışı",
         "ASELSAN'ın ürettiği veya modernize ettiği her insansız araç başına yazılım lisans bedeli faturalandırılması."),
        ("2. İlave Güvenlik Paketleri",
         "Temel sürüşe ek olarak; Mayın Arama Paketi, Gaz Kaçınma Paketi veya Sürü Halinde İlerleme Paketi gibi ilave yazılım modülleri satışı."),
        ("3. Araç Montajı, Saha Testi ve Yıllık Bakım",
         "Yeni araçlara sensör ve yazılım bağlama desteği, arazide test desteği ve yıllık yazılım güncelleme/bakım sözleşmeleri.")
    ]
    for i, (m_title, m_desc) in enumerate(models):
        top_pos = Inches(1.7 + i * 1.6)
        card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.7), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = C_LIGHT
        card.line.color.rgb = C_BLUE
        card.line.width = Pt(1.5)
        tb = s8.shapes.add_textbox(Inches(1.1), top_pos + Inches(0.15), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = m_title
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = C_NAVY
        pd = tf.add_paragraph()
        pd.text = m_desc
        pd.font.size = Pt(12)
        pd.font.color.rgb = C_DARK

    # SLIDE 9: PAZARA GİRİŞ
    s9 = prs.slides[8]
    clear_slide(s9)
    add_header(s9, "PAZARA GİRİŞ: Adım Adım Büyüme Planımız")
    phases = [
        ("1. Aşama: ASELSAN Araçlarında Saha Denemesi (2026)",
         "ASELSAN Girişimcilik Merkezi (Axcelerate) desteğiyle ASELSAN taktik kara araçlarında yazılımımızı arazide denemek ve ilk ortak başarıyı yakalamak."),
        ("2. Aşama: Türk Savunma Sanayii Araçlarına Yayılma (2027)",
         "ASELSAN onaylı tedarikçi gücüyle; yerli zırhlı araç üreticilerinin araçlarına otonom devriye ve konvoy yazılımı sağlamak."),
        ("3. Aşama: Dost Ülkelere İhracat (2027-2028)",
         "NATO tedarikçi kodumuz ve tescillerimizle, dost ülke ordularının projelerine yerli yazılımımızı ihraç etmek.")
    ]
    for i, (p_title, p_desc) in enumerate(phases):
        top_pos = Inches(1.7 + i * 1.6)
        card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.7), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = C_LIGHT
        card.line.color.rgb = C_BORDER
        tb = s9.shapes.add_textbox(Inches(1.1), top_pos + Inches(0.15), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = p_title
        pt.font.size = Pt(15)
        pt.font.bold = True
        pt.font.color.rgb = C_NAVY
        pd = tf.add_paragraph()
        pd.text = p_desc
        pd.font.size = Pt(12)
        pd.font.color.rgb = C_DARK

    # SLIDE 10: YOL HARİTASI
    s10 = prs.slides[9]
    clear_slide(s10)
    add_header(s10, "YOL HARİTASI: 2026 - 2028 Hedeflerimiz")
    roadmap_steps = [
        ("ŞU AN", "Çalışır Sistem", "16 bin satır yerli kod, 1.301 test ile doğrulandı. ASELSAN Tedarikçi Portalı Yazılım Ön Onayı alındı."),
        ("YIL SONU", "Saha Denemesi", "Axcelerate programı ile ASELSAN taktik aracında kapalı alanda ve arazide ilk sürüş denemesinin yapılması."),
        ("2027 İLK YARI", "Seri Lisanslama", "İlk 10 taktik kara aracı için seri yazılım lisans teslimatı ve kontrol ekranı tam uyumu."),
        ("2027 - 2028", "Sürü ve İhracat", "Araçların sürü halinde birlikte hareket etmesi ve dost ülkelere yazılım satışı.")
    ]
    for idx, (q_title, q_sub, q_desc) in enumerate(roadmap_steps):
        l = Inches(0.8 + idx * 3.0)
        card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, Inches(1.8), Inches(2.7), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = C_LIGHT
        card.line.color.rgb = C_BLUE if idx <= 1 else C_BORDER
        card.line.width = Pt(1.5) if idx <= 1 else Pt(1)
        tb = s10.shapes.add_textbox(l + Inches(0.2), Inches(2.0), Inches(2.3), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = q_title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = C_BLUE if idx <= 1 else C_NAVY
        p2 = tf.add_paragraph()
        p2.text = q_sub
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = C_DARK
        p3 = tf.add_paragraph()
        p3.text = q_desc
        p3.font.size = Pt(11)
        p3.font.color.rgb = C_GRAY

    # SLIDE 11: EKİP
    s11 = prs.slides[10]
    clear_slide(s11)
    add_header(s11, "EKİP: Vizyonumuz ve Mühendislik Gücümüz")
    team = [
        ("Murat Furkan Bayram", "Kurucu & CEO / Sistem Mimarı", 
         "• 16.000 satırlık Trustia Otonomi Çekirdeğinin mimarı ve yazarı.\n• 1.301 otomatik test senaryosunun geliştiricisi (%100 başarı).\n• TÜBİTAK ARBİS, BTK Savunma ve KOSGEB İleri Girişimci tescilleri sahibi.\n• Çekirdek Yazılım, Rota Planlama ve Algoritma Lideri."),
        ("Denizcan Özcan", "Baş Donanım ve Entegrasyon Mühendisi", 
         "• ASELSAN Aday Mühendislik Havuzu Üyesi.\n• TEKNOFEST Robotaksi Finalisti.\n• Lazer tarayıcılar (LiDAR), sensörler, gömülü bilgisayarlar ve araç içi haberleşme uzmanı.\n• Donanım Kurulumu ve Saha Testleri Lideri.")
    ]
    for idx, (name, role, bio) in enumerate(team):
        l = Inches(1.2 + idx * 5.6)
        card = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, Inches(1.8), Inches(5.2), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = C_LIGHT
        card.line.color.rgb = C_BORDER
        tb = s11.shapes.add_textbox(l + Inches(0.3), Inches(2.1), Inches(4.6), Inches(4.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = name
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY
        p2 = tf.add_paragraph()
        p2.text = role
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = C_BLUE
        p3 = tf.add_paragraph()
        p3.text = bio
        p3.font.size = Pt(11)
        p3.font.color.rgb = C_DARK

    # SLIDE 12: BEKLENTİLER
    s12 = prs.slides[11]
    clear_slide(s12)
    add_header(s12, "ASELSAN'DAN BEKLENTİLERİMİZ: Birlikte Neler Yapabiliriz?")
    expectations = [
        ("1. ASELSAN Araçlarında Ortak Saha Denemesi",
         "ASELSAN'ın mevcut insansız kara araçlarında yazılımımızı araca yükleyip arazide birlikte denemek."),
        ("2. Gerçekçi Askeri Test Alanı Desteği",
         "GPS uydularının kapalı olduğu kontrollü askeri test alanlarında yazılımımızın gücünü resmi olarak kanıtlamak."),
        ("3. Yerli Yazılım Tedarikçi Anlaşması",
         "Ön onayı tamamlanan Yazılım Geliştirme başvurumuzla birlikte ASELSAN'ın onaylı yazılım tedarikçisi olmak ve araç başı lisans modeliyle çalışmak."),
        ("4. Askeri Standartlar ve Uzman Desteği",
         "Askeri çevre koşullarına ve güvenlik kurallarına uyum sürecinde ASELSAN uzmanlarının tecrübesinden faydalanmak.")
    ]
    for idx, (etitle, edesc) in enumerate(expectations):
        row = idx // 2
        col = idx % 2
        l = Inches(0.8 + col * 6.0)
        t = Inches(1.7 + row * 2.5)
        card = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(5.7), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = C_LIGHT
        card.line.color.rgb = C_BLUE
        card.line.width = Pt(1.5)
        tb = s12.shapes.add_textbox(l + Inches(0.25), t + Inches(0.2), Inches(5.2), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = etitle
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = C_NAVY
        p2 = tf.add_paragraph()
        p2.text = edesc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = C_DARK

    # SLIDE 13: KAPANIŞ
    s13 = prs.slides[12]
    clear_slide(s13)
    bg_card = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.7), Inches(5.9))
    bg_card.fill.solid()
    bg_card.fill.fore_color.rgb = C_LIGHT
    bg_card.line.color.rgb = C_BORDER
    bg_card.line.width = Pt(1.5)

    if os.path.exists(logo_path):
        s13.shapes.add_picture(logo_path, Inches(1.3), Inches(1.3), width=Inches(1.8))

    tb = s13.shapes.add_textbox(Inches(1.3), Inches(2.2), Inches(10.5), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "TEŞEKKÜR EDERİZ"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = C_BLUE

    p1 = tf.add_paragraph()
    p1.text = "Milli Otonomi, Güçlü Savunma."
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = C_NAVY

    p2 = tf.add_paragraph()
    p2.text = "ASELSAN ile birlikte sahada Mehmetçiğin can güvenliğini yerli otonomiyle korumak için hazırız."
    p2.font.size = Pt(14)
    p2.font.color.rgb = C_GRAY

    contact_box = s13.shapes.add_textbox(Inches(1.3), Inches(4.5), Inches(10.5), Inches(1.8))
    ctf = contact_box.text_frame
    ctf.word_wrap = True
    cp1 = ctf.paragraphs[0]
    cp1.text = "İLETİŞİM BİLGİLERİ:"
    cp1.font.size = Pt(12)
    cp1.font.bold = True
    cp1.font.color.rgb = C_NAVY
    cp2 = ctf.add_paragraph()
    cp2.text = "Murat Furkan Bayram | Kurucu & Sistem Mimarı"
    cp2.font.size = Pt(11.5)
    cp2.font.bold = True
    cp2.font.color.rgb = C_DARK
    cp3 = ctf.add_paragraph()
    cp3.text = "E-posta: murat@trustia.com.tr | kariyer@trustia.com.tr  •  Telefon: +90 537 064 04 60"
    cp3.font.size = Pt(11)
    cp3.font.color.rgb = C_DARK
    cp4 = ctf.add_paragraph()
    cp4.text = "Web: trustia.com.tr  •  Merkez: İTO BTM Fulya Kampüsü, Şişli / İstanbul"
    cp4.font.size = Pt(11)
    cp4.font.color.rgb = C_GRAY

    prs.save(output_path_1)
    prs.save(output_path_2)
    print(f"PPTX Updated with ultra-clean Turkish text at {output_path_1} and {output_path_2}")

if __name__ == "__main__":
    build_deck()
